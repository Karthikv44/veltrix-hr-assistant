"""Fixtures and programmatic builders for document formats: PDF, DOCX, Markdown, HTML."""

from pathlib import Path

import pymupdf
from docx import Document


def create_markdown_fixture(
    file_path: Path,
    title: str = "Test Policy",
    sections: list[tuple[str, str]] | None = None,
    table: list[list[str]] | None = None,
) -> None:
    """Create a synthetic Markdown document fixture."""
    lines: list[str] = [f"# {title}\n"]
    if sections:
        for heading, body in sections:
            lines.append(f"## {heading}\n")
            lines.append(f"{body}\n")
    if table:
        lines.append("\n| " + " | ".join(table[0]) + " |")
        lines.append("| " + " | ".join(["---"] * len(table[0])) + " |")
        for row in table[1:]:
            lines.append("| " + " | ".join(row) + " |")
        lines.append("\n")

    file_path.write_text("\n".join(lines), encoding="utf-8")


def create_html_fixture(
    file_path: Path,
    title: str = "Test FAQ",
    sections: list[tuple[str, str]] | None = None,
    table: list[list[str]] | None = None,
) -> None:
    """Create a synthetic HTML document fixture."""
    html_parts: list[str] = [
        "<!DOCTYPE html><html><head><title>Test</title></head><body>",
        f"<h1>{title}</h1>",
    ]
    if sections:
        for heading, body in sections:
            html_parts.append(f"<h2>{heading}</h2>")
            html_parts.append(f"<p>{body}</p>")
    if table:
        html_parts.append("<table>")
        html_parts.append("<tr>" + "".join(f"<th>{c}</th>" for c in table[0]) + "</tr>")
        for row in table[1:]:
            html_parts.append("<tr>" + "".join(f"<td>{c}</td>" for c in row) + "</tr>")
        html_parts.append("</table>")
    html_parts.append("</body></html>")

    file_path.write_text("\n".join(html_parts), encoding="utf-8")


def create_docx_fixture(
    file_path: Path,
    title: str = "Test Docx",
    sections: list[tuple[str, str]] | None = None,
    table: list[list[str]] | None = None,
) -> None:
    """Create a synthetic Word DOCX document fixture."""
    doc = Document()
    doc.add_heading(title, level=0)
    if sections:
        for heading, body in sections:
            doc.add_heading(heading, level=1)
            doc.add_paragraph(body)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for r_idx, row in enumerate(table):
            for c_idx, cell in enumerate(row):
                t.cell(r_idx, c_idx).text = cell

    doc.save(str(file_path))


def create_pdf_fixture(
    file_path: Path,
    title: str = "Test PDF",
    sections: list[tuple[str, str]] | None = None,
) -> None:
    """Create a synthetic PDF document fixture using PyMuPDF."""
    doc = pymupdf.open()
    page = doc.new_page()
    p = pymupdf.Point(50, 50)
    page.insert_text(p, title, fontsize=18)

    y = 80
    if sections:
        for heading, body in sections:
            page.insert_text(pymupdf.Point(50, y), heading, fontsize=14)
            y += 25
            page.insert_text(pymupdf.Point(50, y), body, fontsize=10)
            y += 40

    doc.save(str(file_path))
    doc.close()
